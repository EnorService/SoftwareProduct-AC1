from __future__ import annotations

from collections import Counter
from datetime import datetime
from pathlib import Path
from threading import RLock
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook


class DataError(Exception):
    """Erro de regra ou de consistência dos dados do aplicativo."""


class ExcelDB:
    """Pequena camada de persistência usando um único arquivo Excel.

    O projeto não utiliza SQL. Cada aba funciona como uma coleção de registros
    e os relacionamentos são controlados pelas chaves ID armazenadas nas abas.
    """

    SCHEMAS: dict[str, list[str]] = {
        "Produtos": [
            "id",
            "nome",
            "descricao",
            "categoria",
            "preco",
            "estoque",
            "id_fornecedor",
            "ativo",
            "criado_em",
        ],
        "Clientes": [
            "id",
            "nome",
            "cpf",
            "telefone",
            "email",
            "endereco",
            "ativo",
            "criado_em",
        ],
        "Fornecedores": [
            "id",
            "nome",
            "cnpj",
            "telefone",
            "email",
            "endereco",
            "ativo",
            "criado_em",
        ],
        "Filiais": [
            "id",
            "nome",
            "cnpj",
            "telefone",
            "endereco",
            "cidade",
            "uf",
            "ativo",
            "criado_em",
        ],
        "Pedidos": [
            "id",
            "numero",
            "id_cliente",
            "id_filial",
            "data_pedido",
            "status",
            "total",
            "observacao",
            "efetivado_em",
        ],
        "ItensPedido": [
            "id",
            "id_pedido",
            "id_produto",
            "quantidade",
            "preco_unitario",
            "subtotal",
        ],
        "Configuracoes": [
            "chave",
            "valor",
            "descricao",
        ],
    }

    FEATURE_DEFAULTS: dict[str, tuple[str, str]] = {
        "ac1": ("1", "AC1 · Produtos"),
        "ac2": ("1", "AC2 · Clientes"),
        "ac3": ("1", "AC3 · Fornecedores"),
        "final": ("1", "Prova final · Filiais e pedidos"),
    }

    RELATIONSHIPS: dict[str, tuple[str, str]] = {
        "Produtos": ("ItensPedido", "id_produto"),
        "Clientes": ("Pedidos", "id_cliente"),
        "Fornecedores": ("Produtos", "id_fornecedor"),
        "Filiais": ("Pedidos", "id_filial"),
    }

    def __init__(self, path: str | Path):
        self.path = Path(path)
        self._lock = RLock()
        self.ensure_workbook()

    @staticmethod
    def _as_id(value: Any) -> int | None:
        if value in (None, ""):
            return None
        try:
            return int(value)
        except (TypeError, ValueError):
            return None

    @staticmethod
    def _is_blank_row(values: Iterable[Any]) -> bool:
        return all(value in (None, "") for value in values)

    @staticmethod
    def _is_active(record: dict[str, Any]) -> bool:
        value = record.get("ativo")
        return value is not False and str(value).strip().upper() not in {
            "FALSE",
            "0",
            "INATIVO",
            "N",
        }

    def ensure_workbook(self) -> None:
        with self._lock:
            self.path.parent.mkdir(parents=True, exist_ok=True)

            if self.path.exists():
                workbook = load_workbook(self.path)
            else:
                workbook = Workbook()

            if "Sheet" in workbook.sheetnames and len(workbook.sheetnames) == 1:
                del workbook["Sheet"]

            for sheet_name, headers in self.SCHEMAS.items():
                if sheet_name not in workbook.sheetnames:
                    worksheet = workbook.create_sheet(sheet_name)
                    worksheet.append(headers)
                    continue

                worksheet = workbook[sheet_name]
                current_headers = [
                    str(cell.value).strip().lower()
                    for cell in worksheet[1]
                    if cell.value not in (None, "")
                ]
                if not current_headers:
                    worksheet.append(headers)
                    continue

                for header in headers:
                    if header not in current_headers:
                        worksheet.cell(row=1, column=worksheet.max_column + 1).value = header

            settings_sheet = workbook["Configuracoes"]
            configured_keys = {
                str(row[0].value).strip()
                for row in settings_sheet.iter_rows(min_row=2, max_col=1)
                if row[0].value not in (None, "")
            }
            for feature_key, (default_value, description) in self.FEATURE_DEFAULTS.items():
                setting_key = f"feature.{feature_key}"
                if setting_key not in configured_keys:
                    settings_sheet.append([setting_key, default_value, description])
            settings_sheet.freeze_panes = "A2"
            settings_sheet.auto_filter.ref = settings_sheet.dimensions

            workbook.save(self.path)
            workbook.close()

    def _read_sheet(self, workbook, sheet_name: str) -> list[dict[str, Any]]:
        if sheet_name not in self.SCHEMAS:
            raise DataError(f"A aba '{sheet_name}' não existe.")

        worksheet = workbook[sheet_name]
        headers = [
            str(cell.value).strip().lower() if cell.value not in (None, "") else ""
            for cell in worksheet[1]
        ]
        rows: list[dict[str, Any]] = []
        for values in worksheet.iter_rows(min_row=2, values_only=True):
            if self._is_blank_row(values):
                continue
            row = {
                header: values[index] if index < len(values) else None
                for index, header in enumerate(headers)
                if header
            }
            for field in self.SCHEMAS[sheet_name]:
                row.setdefault(field, None)
            rows.append(row)
        return rows

    def _write_sheet(self, workbook, sheet_name: str, rows: list[dict[str, Any]]) -> None:
        worksheet = workbook[sheet_name]
        worksheet.delete_rows(1, worksheet.max_row)
        headers = self.SCHEMAS[sheet_name]
        worksheet.append(headers)
        for row in rows:
            worksheet.append([row.get(header) for header in headers])
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions

        widths = {
            "id": 10,
            "nome": 28,
            "descricao": 36,
            "endereco": 36,
            "observacao": 40,
        }
        for column_index, header in enumerate(headers, start=1):
            worksheet.column_dimensions[chr(64 + column_index) if column_index <= 26 else "A"].width = widths.get(header, 18)

    def _save_workbook(self, workbook) -> None:
        workbook.save(self.path)
        workbook.close()

    def list_records(self, sheet_name: str, active_only: bool = False) -> list[dict[str, Any]]:
        with self._lock:
            workbook = load_workbook(self.path)
            try:
                records = self._read_sheet(workbook, sheet_name)
                if active_only:
                    records = [record for record in records if self._is_active(record)]
                return records
            finally:
                workbook.close()

    def feature_flags(self) -> dict[str, bool]:
        settings = {
            str(row.get("chave")): str(row.get("valor") or "1")
            for row in self.list_records("Configuracoes")
        }
        return {
            feature_key: settings.get(f"feature.{feature_key}", default_value) == "1"
            for feature_key, (default_value, _description) in self.FEATURE_DEFAULTS.items()
        }

    def feature_enabled(self, feature_key: str) -> bool:
        return self.feature_flags().get(feature_key, True)

    def get_setting(self, key: str, default: Any = None) -> Any:
        for row in self.list_records("Configuracoes"):
            if str(row.get("chave") or "").strip() == key:
                return row.get("valor")
        return default

    def set_setting(self, key: str, value: Any, description: str = "") -> None:
        with self._lock:
            workbook = load_workbook(self.path)
            try:
                rows = self._read_sheet(workbook, "Configuracoes")
                row = next(
                    (item for item in rows if str(item.get("chave") or "").strip() == key),
                    None,
                )
                if row is None:
                    rows.append({"chave": key, "valor": value, "descricao": description})
                else:
                    row["valor"] = value
                    if description:
                        row["descricao"] = description
                self._write_sheet(workbook, "Configuracoes", rows)
                self._save_workbook(workbook)
            except Exception:
                workbook.close()
                raise

    def set_feature_flags(self, flags: dict[str, bool]) -> None:
        with self._lock:
            workbook = load_workbook(self.path)
            try:
                rows = self._read_sheet(workbook, "Configuracoes")
                rows_by_key = {str(row.get("chave")): row for row in rows}
                for feature_key, (default_value, description) in self.FEATURE_DEFAULTS.items():
                    setting_key = f"feature.{feature_key}"
                    row = rows_by_key.get(setting_key)
                    if row is None:
                        row = {"chave": setting_key, "valor": default_value, "descricao": description}
                        rows.append(row)
                    row["valor"] = "1" if flags.get(feature_key, False) else "0"
                    row["descricao"] = description
                self._write_sheet(workbook, "Configuracoes", rows)
                self._save_workbook(workbook)
            except Exception:
                workbook.close()
                raise

    def get(self, sheet_name: str, record_id: int) -> dict[str, Any] | None:
        record_id = int(record_id)
        return next(
            (
                record
                for record in self.list_records(sheet_name)
                if self._as_id(record.get("id")) == record_id
            ),
            None,
        )

    def _next_id(self, rows: list[dict[str, Any]]) -> int:
        ids = [self._as_id(row.get("id")) or 0 for row in rows]
        return max(ids, default=0) + 1

    def insert(self, sheet_name: str, values: dict[str, Any]) -> dict[str, Any]:
        with self._lock:
            workbook = load_workbook(self.path)
            try:
                rows = self._read_sheet(workbook, sheet_name)
                record = {field: values.get(field) for field in self.SCHEMAS[sheet_name]}
                record["id"] = self._next_id(rows)
                rows.append(record)
                self._write_sheet(workbook, sheet_name, rows)
                self._save_workbook(workbook)
                return record
            except Exception:
                workbook.close()
                raise

    def update(self, sheet_name: str, record_id: int, values: dict[str, Any]) -> dict[str, Any]:
        with self._lock:
            workbook = load_workbook(self.path)
            try:
                rows = self._read_sheet(workbook, sheet_name)
                record = next(
                    (
                        row
                        for row in rows
                        if self._as_id(row.get("id")) == int(record_id)
                    ),
                    None,
                )
                if record is None:
                    raise DataError("Registro não encontrado.")
                for field, value in values.items():
                    if field in self.SCHEMAS[sheet_name] and field != "id":
                        record[field] = value
                self._write_sheet(workbook, sheet_name, rows)
                self._save_workbook(workbook)
                return record
            except Exception:
                workbook.close()
                raise

    def delete_or_deactivate(self, sheet_name: str, record_id: int) -> str:
        """Exclui fisicamente quando possível; inativa se houver dependências."""
        relation = self.RELATIONSHIPS.get(sheet_name)
        if relation:
            target_sheet, target_field = relation
            references = [
                row
                for row in self.list_records(target_sheet)
                if self._as_id(row.get(target_field)) == int(record_id)
            ]
            if references:
                self.update(sheet_name, record_id, {"ativo": False})
                return "inativado"

        with self._lock:
            workbook = load_workbook(self.path)
            try:
                rows = self._read_sheet(workbook, sheet_name)
                filtered = [
                    row
                    for row in rows
                    if self._as_id(row.get("id")) != int(record_id)
                ]
                if len(filtered) == len(rows):
                    raise DataError("Registro não encontrado.")
                self._write_sheet(workbook, sheet_name, filtered)
                self._save_workbook(workbook)
                return "excluido"
            except Exception:
                workbook.close()
                raise

    def create_order(
        self,
        client_id: int,
        branch_id: int,
        items: list[tuple[int, int]],
        observation: str = "",
    ) -> dict[str, Any]:
        with self._lock:
            workbook = load_workbook(self.path)
            try:
                clients = self._read_sheet(workbook, "Clientes")
                branches = self._read_sheet(workbook, "Filiais")
                products = self._read_sheet(workbook, "Produtos")
                orders = self._read_sheet(workbook, "Pedidos")
                order_items = self._read_sheet(workbook, "ItensPedido")

                client = next(
                    (row for row in clients if self._as_id(row.get("id")) == int(client_id)),
                    None,
                )
                branch = next(
                    (row for row in branches if self._as_id(row.get("id")) == int(branch_id)),
                    None,
                )
                if not client or not self._is_active(client):
                    raise DataError("Selecione um cliente ativo.")
                if not branch or not self._is_active(branch):
                    raise DataError("Selecione uma filial ativa.")

                merged: dict[int, int] = {}
                for product_id, quantity in items:
                    product_id = int(product_id)
                    quantity = int(quantity)
                    if quantity <= 0:
                        raise DataError("A quantidade dos produtos deve ser maior que zero.")
                    merged[product_id] = merged.get(product_id, 0) + quantity

                if not merged:
                    raise DataError("Adicione pelo menos um produto ao pedido.")

                product_map = {
                    self._as_id(product.get("id")): product for product in products
                }
                order_id = self._next_id(orders)
                total = 0.0
                new_items: list[dict[str, Any]] = []
                next_item_id = self._next_id(order_items)

                for product_id, quantity in merged.items():
                    product = product_map.get(product_id)
                    if not product or not self._is_active(product):
                        raise DataError("Um dos produtos selecionados está inativo ou não existe.")
                    price = float(product.get("preco") or 0)
                    subtotal = round(price * quantity, 2)
                    total += subtotal
                    new_items.append(
                        {
                            "id": next_item_id,
                            "id_pedido": order_id,
                            "id_produto": product_id,
                            "quantidade": quantity,
                            "preco_unitario": price,
                            "subtotal": subtotal,
                        }
                    )
                    next_item_id += 1

                order = {
                    "id": order_id,
                    "numero": f"PED-{order_id:05d}",
                    "id_cliente": int(client_id),
                    "id_filial": int(branch_id),
                    "data_pedido": datetime.now().isoformat(timespec="minutes"),
                    "status": "ABERTO",
                    "total": round(total, 2),
                    "observacao": observation.strip(),
                    "efetivado_em": None,
                }
                orders.append(order)
                order_items.extend(new_items)
                self._write_sheet(workbook, "Pedidos", orders)
                self._write_sheet(workbook, "ItensPedido", order_items)
                self._save_workbook(workbook)
                return order
            except Exception:
                workbook.close()
                raise

    def finalize_order(self, order_id: int) -> dict[str, Any]:
        with self._lock:
            workbook = load_workbook(self.path)
            try:
                orders = self._read_sheet(workbook, "Pedidos")
                order_items = self._read_sheet(workbook, "ItensPedido")
                products = self._read_sheet(workbook, "Produtos")

                order = next(
                    (row for row in orders if self._as_id(row.get("id")) == int(order_id)),
                    None,
                )
                if order is None:
                    raise DataError("Pedido não encontrado.")
                if str(order.get("status")).upper() != "ABERTO":
                    raise DataError("Somente pedidos abertos podem ser efetivados.")

                items = [
                    item
                    for item in order_items
                    if self._as_id(item.get("id_pedido")) == int(order_id)
                ]
                product_map = {
                    self._as_id(product.get("id")): product for product in products
                }
                required: Counter[int] = Counter()
                for item in items:
                    product_id = self._as_id(item.get("id_produto"))
                    required[product_id] += int(item.get("quantidade") or 0)

                for product_id, quantity in required.items():
                    product = product_map.get(product_id)
                    if not product:
                        raise DataError("Um produto do pedido não foi encontrado.")
                    stock = int(product.get("estoque") or 0)
                    if stock < quantity:
                        raise DataError(
                            f"Estoque insuficiente para '{product.get('nome')}'. Disponível: {stock}."
                        )

                for product_id, quantity in required.items():
                    product_map[product_id]["estoque"] = int(product_map[product_id].get("estoque") or 0) - quantity

                order["status"] = "CONFIRMADO"
                order["efetivado_em"] = datetime.now().isoformat(timespec="minutes")
                self._write_sheet(workbook, "Produtos", products)
                self._write_sheet(workbook, "Pedidos", orders)
                self._save_workbook(workbook)
                return order
            except Exception:
                workbook.close()
                raise

    def cancel_order(self, order_id: int) -> dict[str, Any]:
        order = self.get("Pedidos", order_id)
        if not order:
            raise DataError("Pedido não encontrado.")
        if str(order.get("status")).upper() != "ABERTO":
            raise DataError("Somente pedidos abertos podem ser cancelados.")
        return self.update("Pedidos", order_id, {"status": "CANCELADO"})

    def order_details(self, order_id: int) -> dict[str, Any] | None:
        order = self.get("Pedidos", order_id)
        if not order:
            return None
        client = self.get("Clientes", self._as_id(order.get("id_cliente")))
        branch = self.get("Filiais", self._as_id(order.get("id_filial")))
        items = []
        for item in self.list_records("ItensPedido"):
            if self._as_id(item.get("id_pedido")) != int(order_id):
                continue
            product = self.get("Produtos", self._as_id(item.get("id_produto"))) or {}
            items.append({**item, "produto_nome": product.get("nome", "Produto removido")})
        return {"order": order, "client": client, "branch": branch, "items": items}

    def enriched_orders(self) -> list[dict[str, Any]]:
        clients = {self._as_id(row.get("id")): row for row in self.list_records("Clientes")}
        branches = {self._as_id(row.get("id")): row for row in self.list_records("Filiais")}
        result = []
        for order in self.list_records("Pedidos"):
            result.append(
                {
                    **order,
                    "cliente_nome": (clients.get(self._as_id(order.get("id_cliente"))) or {}).get("nome", "-"),
                    "filial_nome": (branches.get(self._as_id(order.get("id_filial"))) or {}).get("nome", "-"),
                }
            )
        return sorted(result, key=lambda row: self._as_id(row.get("id")) or 0, reverse=True)

    def dashboard(self) -> dict[str, Any]:
        sheets = ["Produtos", "Clientes", "Fornecedores", "Filiais"]
        counts = {
            sheet: len(self.list_records(sheet, active_only=True)) for sheet in sheets
        }
        orders = self.list_records("Pedidos")
        statuses = Counter(str(order.get("status") or "ABERTO").upper() for order in orders)
        total_sales = sum(
            float(order.get("total") or 0)
            for order in orders
            if str(order.get("status")).upper() == "CONFIRMADO"
        )
        return {
            "counts": counts,
            "orders_total": len(orders),
            "orders_status": dict(statuses),
            "confirmed_sales": round(total_sales, 2),
            "recent_orders": self.enriched_orders()[:5],
        }
